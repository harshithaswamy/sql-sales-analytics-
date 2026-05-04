"""
SQL Sales Analytics Dashboard — Data Export Script
Author: Harshitha S
Description: Connects to MySQL, runs KPI queries, exports results to Excel with
             multiple sheets and basic formatting.
"""

import mysql.connector
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import os

# ─── CONFIG ───────────────────────────────────────────────────
DB_CONFIG = {
    "host":     "localhost",
    "user":     "root",          # Change to your MySQL username
    "password": "your_password", # Change to your MySQL password
    "database": "sales_analytics"
}

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), '..', 'excel', 'reports')
os.makedirs(OUTPUT_DIR, exist_ok=True)

# ─── QUERIES ──────────────────────────────────────────────────
QUERIES = {
    "Monthly Revenue": """
        SELECT DATE_FORMAT(order_date,'%Y-%m') AS month,
               COUNT(order_id) AS total_orders,
               ROUND(SUM(amount),2) AS total_revenue
        FROM orders
        GROUP BY month ORDER BY month
    """,
    "MoM Growth": """
        WITH m AS (
            SELECT DATE_FORMAT(order_date,'%Y-%m') AS month,
                   ROUND(SUM(amount),2) AS revenue
            FROM orders GROUP BY month
        )
        SELECT month, revenue,
               LAG(revenue) OVER (ORDER BY month) AS prev_revenue,
               ROUND((revenue - LAG(revenue) OVER (ORDER BY month))
                     / NULLIF(LAG(revenue) OVER (ORDER BY month),0)*100,2) AS mom_growth_pct
        FROM m ORDER BY month
    """,
    "Top Products": """
        SELECT p.product_name, p.category,
               COUNT(o.order_id) AS times_ordered,
               SUM(o.quantity) AS units_sold,
               ROUND(SUM(o.amount),2) AS total_revenue
        FROM orders o JOIN products p ON o.product_id=p.product_id
        GROUP BY p.product_id, p.product_name, p.category
        ORDER BY total_revenue DESC
    """,
    "Revenue by Region": """
        SELECT r.region_name,
               COUNT(o.order_id) AS total_orders,
               ROUND(SUM(o.amount),2) AS total_revenue
        FROM orders o JOIN regions r ON o.region_id=r.region_id
        GROUP BY r.region_name ORDER BY total_revenue DESC
    """,
    "Customer Segments": """
        SELECT c.segment,
               COUNT(DISTINCT o.customer_id) AS unique_customers,
               COUNT(o.order_id) AS total_orders,
               ROUND(SUM(o.amount),2) AS total_revenue
        FROM orders o JOIN customers c ON o.customer_id=c.customer_id
        GROUP BY c.segment ORDER BY total_revenue DESC
    """,
    "Return Analysis": """
        SELECT reason, COUNT(*) AS return_count,
               ROUND(COUNT(*) / (SELECT COUNT(*) FROM returns)*100,2) AS pct
        FROM returns GROUP BY reason ORDER BY return_count DESC
    """
}

# ─── HELPERS ──────────────────────────────────────────────────
HEADER_FILL  = PatternFill("solid", fgColor="1F5C8B")
HEADER_FONT  = Font(bold=True, color="FFFFFF", name="Arial", size=10)
ALT_FILL     = PatternFill("solid", fgColor="E8F0F7")
THIN_BORDER  = Border(
    bottom=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC")
)

def style_sheet(ws):
    for col_idx, cell in enumerate(ws[1], 1):
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), 2):
        fill = ALT_FILL if row_idx % 2 == 0 else None
        for cell in row:
            if fill:
                cell.fill = fill
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="left")
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)
    ws.freeze_panes = "A2"

# ─── MAIN ─────────────────────────────────────────────────────
def export():
    print("Connecting to MySQL...")
    conn = mysql.connector.connect(**DB_CONFIG)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = os.path.join(OUTPUT_DIR, f"sales_report_{timestamp}.xlsx")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, query in QUERIES.items():
            print(f"  Exporting: {sheet_name}")
            df = pd.read_sql(query, conn)
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            style_sheet(writer.sheets[sheet_name])

    conn.close()
    print(f"\nExport complete: {output_path}")

if __name__ == "__main__":
    export()
